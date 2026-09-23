from pathlib import Path
from datetime import datetime
import traceback

import pandas as pd

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ==========================================================
# CONFIGURAÇÕES DO LOG
# ==========================================================

COLUNAS_LOG = [
    "ID_Email",
    "Fornecedor",
    "EmailFornecedor",
    "EmailDestino",
    "Assunto",
    "QuantidadeItens",
    "DataTentativa",
    "DataEnvio",
    "DataLimiteRetorno",
    "StatusEnvio",
    "StatusRetorno",
    "DataResposta",
    "EmailResposta",
    "AssuntoResposta",
    "ArquivoResposta",
    "MotivoResposta",
    "ResponsavelAcompanhamento",
    "MotivoObservacao",
    "ArquivoEnviado",
]


COLUNAS_DATA_HORA = [
    "DataTentativa",
    "DataEnvio",
    "DataResposta",
]


COLUNAS_SOMENTE_DATA = [
    "DataLimiteRetorno",
]


STATUS_CONSIDERADOS_RESPONDIDOS = [
    "RESPONDIDO",
    "RESPONDIDO SEM ANEXO",
    "RESPOSTA COM ANEXO INVÁLIDO",
]


# ==========================================================
# FUNÇÕES AUXILIARES
# ==========================================================

def garantir_colunas_log(dataframe):
    """
    Garante que todas as colunas utilizadas no histórico
    existam no DataFrame.

    As colunas existentes são preservadas.
    """

    dataframe = dataframe.copy()

    colunas_texto = [
        "ID_Email",
        "Fornecedor",
        "EmailFornecedor",
        "EmailDestino",
        "Assunto",
        "StatusEnvio",
        "StatusRetorno",
        "EmailResposta",
        "AssuntoResposta",
        "ArquivoResposta",
        "MotivoResposta",
        "ResponsavelAcompanhamento",
        "MotivoObservacao",
        "ArquivoEnviado",
    ]

    colunas_data = [
        "DataTentativa",
        "DataEnvio",
        "DataLimiteRetorno",
        "DataResposta",
    ]

    for coluna in colunas_texto:
        if coluna not in dataframe.columns:
            dataframe[coluna] = ""

    for coluna in colunas_data:
        if coluna not in dataframe.columns:
            dataframe[coluna] = pd.NaT

    if "QuantidadeItens" not in dataframe.columns:
        dataframe["QuantidadeItens"] = 0

    colunas_existentes = [
        coluna
        for coluna in COLUNAS_LOG
        if coluna in dataframe.columns
    ]

    colunas_adicionais = [
        coluna
        for coluna in dataframe.columns
        if coluna not in COLUNAS_LOG
    ]

    return dataframe[
        colunas_existentes
        + colunas_adicionais
    ]


def converter_colunas_data(dataframe):
    """
    Converte as colunas de data para datetime sem timezone.

    O Outlook pode fornecer datas com timezone, porém o Excel
    aceita apenas datas sem timezone.
    """

    dataframe = dataframe.copy()

    for coluna_data in [
        "DataTentativa",
        "DataEnvio",
        "DataLimiteRetorno",
        "DataResposta",
    ]:
        if coluna_data not in dataframe.columns:
            continue

        serie_data = pd.to_datetime(
            dataframe[coluna_data],
            dayfirst=True,
            errors="coerce",
            utc=True,
        )

        dataframe[coluna_data] = (
            serie_data.dt.tz_localize(None)
            .astype("datetime64[ns]")
        )

    return dataframe


def normalizar_status(
    dataframe,
    nome_coluna,
):
    """
    Retorna uma Series com os valores normalizados
    de uma coluna de status.
    """

    if nome_coluna not in dataframe.columns:
        return pd.Series(
            "",
            index=dataframe.index,
            dtype="string",
        )

    return (
        dataframe[nome_coluna]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )


def ordenar_dataframe_log(dataframe):
    """
    Ordena o DataFrame pelas datas disponíveis.

    Prioridade:
    1. DataEnvio;
    2. DataTentativa.
    """

    dataframe = dataframe.copy()

    colunas_ordenacao = []

    if "DataEnvio" in dataframe.columns:
        colunas_ordenacao.append(
            "DataEnvio"
        )

    if "DataTentativa" in dataframe.columns:
        colunas_ordenacao.append(
            "DataTentativa"
        )

    if (
        not dataframe.empty
        and colunas_ordenacao
    ):
        dataframe.sort_values(
            by=colunas_ordenacao,
            ascending=[
                False
            ] * len(colunas_ordenacao),
            na_position="last",
            inplace=True,
        )

    dataframe.reset_index(
        drop=True,
        inplace=True,
    )

    return dataframe


def gerar_visoes_historico(
    df_historico,
):
    """
    Gera todas as abas derivadas usando a aba Historico
    como fonte oficial.

    Retorna um dicionário contendo todos os DataFrames.
    """

    df_historico = garantir_colunas_log(
        df_historico
    )

    df_historico = converter_colunas_data(
        df_historico
    )

    status_envio = normalizar_status(
        dataframe=df_historico,
        nome_coluna="StatusEnvio",
    )

    status_retorno = normalizar_status(
        dataframe=df_historico,
        nome_coluna="StatusRetorno",
    )

    df_enviados = df_historico[
        status_envio.eq("ENVIADO")
    ].copy()

    df_nao_enviados = df_historico[
        status_envio.eq("NÃO ENVIADO")
    ].copy()

    df_aberto_revisao = df_historico[
        status_envio.eq(
            "ABERTO PARA REVISÃO"
        )
    ].copy()

    df_aguardando_retorno = df_historico[
        status_envio.eq("ENVIADO")
        & status_retorno.eq(
            "AGUARDANDO RETORNO"
        )
    ].copy()

    df_respondidos = df_historico[
        status_retorno.isin(
            STATUS_CONSIDERADOS_RESPONDIDOS
        )
    ].copy()

    visoes = {
        "Historico": df_historico,
        "Enviados": df_enviados,
        "Nao_Enviados": df_nao_enviados,
        "Aguardando_Retorno": (
            df_aguardando_retorno
        ),
        "Respondidos": df_respondidos,
        "Aberto_Revisao": df_aberto_revisao,
    }

    for nome_aba, dataframe in visoes.items():
        visoes[nome_aba] = ordenar_dataframe_log(
            dataframe
        )

    return visoes


# ==========================================================
# REGISTRO DO LOG
# ==========================================================

def registrar_log(
    lista_logs,
    fornecedor,
    email_fornecedor,
    email_destino,
    id_email,
    assunto,
    status_envio,
    motivo,
    quantidade_itens,
    caminho_arquivo="",
    data_tentativa=None,
    data_envio=None,
    data_retorno=None,
):
    """
    Registra o resultado do processamento do fornecedor.

    O ID_Email é o mesmo identificador utilizado no assunto
    da mensagem e pode ser pesquisado no Outlook.
    """

    status_envio = str(
        status_envio
    ).strip().upper()

    if data_tentativa is not None:
        data_tentativa = pd.Timestamp(
            data_tentativa
        )

        if data_tentativa.tzinfo is not None:
            data_tentativa = (
                data_tentativa.tz_localize(None)
            )

    else:
        data_tentativa = pd.NaT

    if data_envio is not None:
        data_envio = pd.Timestamp(
            data_envio
        )

        if data_envio.tzinfo is not None:
            data_envio = (
                data_envio.tz_localize(None)
            )

    else:
        data_envio = pd.NaT

    if data_retorno is not None:
        data_retorno = pd.Timestamp(
            data_retorno
        )

        if data_retorno.tzinfo is not None:
            data_retorno = (
                data_retorno.tz_localize(None)
            )

    else:
        data_retorno = pd.NaT

    if status_envio == "ENVIADO":
        status_retorno = (
            "AGUARDANDO RETORNO"
        )

    else:
        status_retorno = ""

    lista_logs.append(
        {
            "ID_Email": str(
                id_email
            ).strip(),
            "Fornecedor": str(
                fornecedor
            ),
            "EmailFornecedor": (
                str(email_fornecedor)
                if email_fornecedor is not None
                else ""
            ),
            "EmailDestino": (
                str(email_destino)
                if email_destino is not None
                else ""
            ),
            "Assunto": str(
                assunto
            ),
            "QuantidadeItens": quantidade_itens,
            "DataTentativa": data_tentativa,
            "DataEnvio": data_envio,
            "DataLimiteRetorno": data_retorno,
            "StatusEnvio": status_envio,
            "StatusRetorno": status_retorno,
            "DataResposta": pd.NaT,
            "EmailResposta": "",
            "AssuntoResposta": "",
            "ArquivoResposta": "",
            "MotivoResposta": "",
            "ResponsavelAcompanhamento": "",
            "MotivoObservacao": str(
                motivo
            ),
            "ArquivoEnviado": (
                str(caminho_arquivo)
                if caminho_arquivo
                else ""
            ),
        }
    )


# ==========================================================
# FORMATAÇÃO DO EXCEL
# ==========================================================

def ajustar_planilha_log(aba):
    """
    Aplica formatação profissional às abas do Excel
    de acompanhamento.
    """

    cor_cabecalho = "FFED1C24"
    cor_fonte_cabecalho = "FFFFFFFF"

    if (
        aba.max_row < 1
        or aba.max_column < 1
    ):
        return

    # ------------------------------------------------------
    # CONGELA O CABEÇALHO E ATIVA O FILTRO
    # ------------------------------------------------------

    aba.freeze_panes = "A2"

    if aba.max_column > 0:
        aba.auto_filter.ref = aba.dimensions

    # ------------------------------------------------------
    # OBTÉM E FORMATA O CABEÇALHO
    # ------------------------------------------------------

    linha_cabecalho = next(
        aba.iter_rows(
            min_row=1,
            max_row=1,
        )
    )

    for celula in linha_cabecalho:
        celula.fill = PatternFill(
            fill_type="solid",
            start_color=cor_cabecalho,
            end_color=cor_cabecalho,
        )

        celula.font = Font(
            color=cor_fonte_cabecalho,
            bold=True,
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    aba.row_dimensions[1].height = 32

    # ------------------------------------------------------
    # IDENTIFICA AS COLUNAS PELOS NOMES
    # ------------------------------------------------------

    nomes_colunas = {
        str(celula.value).strip(): (
            celula.column
        )
        for celula in linha_cabecalho
        if celula.value is not None
    }

    # ------------------------------------------------------
    # AJUSTA A LARGURA E O ALINHAMENTO
    # ------------------------------------------------------

    for coluna in aba.iter_cols():
        if not coluna:
            continue

        comprimento_maximo = 0

        numero_coluna = coluna[0].column

        letra_coluna = get_column_letter(
            numero_coluna
        )

        for celula in coluna:
            if celula.value is not None:
                comprimento = len(
                    str(celula.value)
                )

                comprimento_maximo = max(
                    comprimento_maximo,
                    comprimento,
                )

            if celula.row > 1:
                celula.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        largura = min(
            max(
                comprimento_maximo + 2,
                12,
            ),
            50,
        )

        aba.column_dimensions[
            letra_coluna
        ].width = largura

    # ------------------------------------------------------
    # LARGURAS ESPECÍFICAS
    # ------------------------------------------------------

    larguras_especificas = {
        "ID_Email": 14,
        "Fornecedor": 35,
        "EmailFornecedor": 40,
        "EmailDestino": 40,
        "Assunto": 45,
        "QuantidadeItens": 16,
        "DataTentativa": 20,
        "DataEnvio": 20,
        "DataLimiteRetorno": 20,
        "StatusEnvio": 22,
        "StatusRetorno": 25,
        "DataResposta": 20,
        "EmailResposta": 40,
        "AssuntoResposta": 45,
        "ArquivoResposta": 50,
        "MotivoResposta": 50,
        "ResponsavelAcompanhamento": 30,
        "MotivoObservacao": 50,
        "ArquivoEnviado": 50,
    }

    for nome_coluna, largura in (
        larguras_especificas.items()
    ):
        if nome_coluna not in nomes_colunas:
            continue

        numero_coluna = nomes_colunas[
            nome_coluna
        ]

        letra_coluna = get_column_letter(
            numero_coluna
        )

        aba.column_dimensions[
            letra_coluna
        ].width = largura

    # ------------------------------------------------------
    # FORMATA DATA E HORA
    # ------------------------------------------------------

    for nome_coluna in COLUNAS_DATA_HORA:
        if nome_coluna not in nomes_colunas:
            continue

        numero_coluna = nomes_colunas[
            nome_coluna
        ]

        for numero_linha in range(
            2,
            aba.max_row + 1,
        ):
            aba.cell(
                row=numero_linha,
                column=numero_coluna,
            ).number_format = (
                "dd/mm/yyyy hh:mm"
            )

    # ------------------------------------------------------
    # FORMATA SOMENTE DATA
    # ------------------------------------------------------

    for nome_coluna in COLUNAS_SOMENTE_DATA:
        if nome_coluna not in nomes_colunas:
            continue

        numero_coluna = nomes_colunas[
            nome_coluna
        ]

        for numero_linha in range(
            2,
            aba.max_row + 1,
        ):
            aba.cell(
                row=numero_linha,
                column=numero_coluna,
            ).number_format = (
                "dd/mm/yyyy"
            )

    # ------------------------------------------------------
    # CENTRALIZA COLUNAS ESPECÍFICAS
    # ------------------------------------------------------

    colunas_centralizadas = [
        "ID_Email",
        "QuantidadeItens",
        "DataTentativa",
        "DataEnvio",
        "DataLimiteRetorno",
        "StatusEnvio",
        "StatusRetorno",
        "DataResposta",
    ]

    for nome_coluna in colunas_centralizadas:
        if nome_coluna not in nomes_colunas:
            continue

        numero_coluna = nomes_colunas[
            nome_coluna
        ]

        for numero_linha in range(
            2,
            aba.max_row + 1,
        ):
            aba.cell(
                row=numero_linha,
                column=numero_coluna,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )


# ==========================================================
# SALVAMENTO CENTRALIZADO DAS ABAS
# ==========================================================

def salvar_visoes_log(
    arquivo_log,
    visoes,
):
    """
    Salva o histórico e todas as abas derivadas.

    O arquivo é recriado integralmente a partir da aba
    Historico, considerada a fonte oficial.
    """

    arquivo_log = Path(
        arquivo_log
    )

    arquivo_log.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    ordem_abas = [
        "Historico",
        "Enviados",
        "Nao_Enviados",
        "Aguardando_Retorno",
        "Respondidos",
        "Aberto_Revisao",
    ]

    try:
        with pd.ExcelWriter(
            arquivo_log,
            engine="openpyxl",
            mode="w",
            datetime_format="dd/mm/yyyy hh:mm",
            date_format="dd/mm/yyyy",
        ) as writer:
            for nome_aba in ordem_abas:
                dataframe = visoes.get(
                    nome_aba
                )

                if dataframe is None:
                    dataframe = pd.DataFrame(
                        columns=COLUNAS_LOG
                    )

                dataframe = garantir_colunas_log(
                    dataframe
                )

                dataframe = converter_colunas_data(
                    dataframe
                )

                dataframe.to_excel(
                    writer,
                    sheet_name=nome_aba,
                    index=False,
                )

            for nome_aba in ordem_abas:
                if (
                    nome_aba
                    not in writer.book.sheetnames
                ):
                    continue

                ajustar_planilha_log(
                    writer.book[nome_aba]
                )

    except PermissionError as erro:
        raise PermissionError(
            "Não foi possível atualizar o arquivo "
            "de acompanhamento. Feche o arquivo no Excel "
            "e tente novamente."
        ) from erro

    return True


# ==========================================================
# SALVAMENTO DOS NOVOS LOGS
# ==========================================================

def salvar_log_excel(
    lista_logs,
    arquivo_log,
):
    """
    Salva o histórico de acompanhamento em Excel.

    Abas criadas:
    - Historico;
    - Enviados;
    - Nao_Enviados;
    - Aguardando_Retorno;
    - Respondidos;
    - Aberto_Revisao.

    O arquivo mantém o histórico acumulado.

    O resumo exibido considera somente os registros
    da execução atual.
    """

    if not lista_logs:
        print(
            "Nenhum registro de log para salvar."
        )

        return False

    arquivo_log = Path(
        arquivo_log
    )

    arquivo_log.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ------------------------------------------------------
    # REGISTROS DA EXECUÇÃO ATUAL
    # ------------------------------------------------------

    df_novos_logs = pd.DataFrame(
        lista_logs
    )

    df_novos_logs = garantir_colunas_log(
        df_novos_logs
    )

    df_novos_logs = converter_colunas_data(
        df_novos_logs
    )

    # ------------------------------------------------------
    # RESUMO DA EXECUÇÃO ATUAL
    # ------------------------------------------------------

    status_execucao_atual = normalizar_status(
        dataframe=df_novos_logs,
        nome_coluna="StatusEnvio",
    )

    status_retorno_execucao_atual = (
        normalizar_status(
            dataframe=df_novos_logs,
            nome_coluna="StatusRetorno",
        )
    )

    qtd_processados_execucao = len(
        df_novos_logs
    )

    qtd_enviados_execucao = int(
        status_execucao_atual
        .eq("ENVIADO")
        .sum()
    )

    qtd_nao_enviados_execucao = int(
        status_execucao_atual
        .eq("NÃO ENVIADO")
        .sum()
    )

    qtd_revisao_execucao = int(
        status_execucao_atual
        .eq("ABERTO PARA REVISÃO")
        .sum()
    )

    qtd_aguardando_execucao = int(
        (
            status_execucao_atual.eq(
                "ENVIADO"
            )
            & status_retorno_execucao_atual.eq(
                "AGUARDANDO RETORNO"
            )
        ).sum()
    )

    # ------------------------------------------------------
    # CARREGA O HISTÓRICO ANTERIOR
    # ------------------------------------------------------

    if arquivo_log.exists():
        try:
            df_historico_anterior = pd.read_excel(
                arquivo_log,
                sheet_name="Historico",
                engine="openpyxl",
            )

            df_historico_anterior = (
                garantir_colunas_log(
                    df_historico_anterior
                )
            )

            df_historico_anterior = (
                converter_colunas_data(
                    df_historico_anterior
                )
            )

        except PermissionError as erro:
            raise PermissionError(
                "Não foi possível ler o histórico. "
                "Feche o arquivo no Excel e tente novamente."
            ) from erro

        except ValueError:
            print(
                "A aba Historico não foi encontrada. "
                "Um novo histórico será criado."
            )

            df_historico_anterior = pd.DataFrame(
                columns=COLUNAS_LOG
            )

        except Exception as erro:
            print(
                "Não foi possível ler o histórico anterior: "
                f"{erro}"
            )

            traceback.print_exc()

            df_historico_anterior = pd.DataFrame(
                columns=COLUNAS_LOG
            )

    else:
        df_historico_anterior = pd.DataFrame(
            columns=COLUNAS_LOG
        )

    # ------------------------------------------------------
    # JUNTA HISTÓRICO E NOVOS REGISTROS
    # ------------------------------------------------------

    df_historico = pd.concat(
        [
            df_historico_anterior,
            df_novos_logs,
        ],
        ignore_index=True,
    )

    df_historico = garantir_colunas_log(
        df_historico
    )

    df_historico = converter_colunas_data(
        df_historico
    )

    # Um ID_Email representa uma tentativa de envio.
    # Quando o mesmo ID for atualizado, mantém a versão
    # mais recente do registro.
    if "ID_Email" in df_historico.columns:
        ids_validos = (
            df_historico["ID_Email"]
            .fillna("")
            .astype(str)
            .str.strip()
            .ne("")
        )

        df_com_id = (
            df_historico[
                ids_validos
            ]
            .drop_duplicates(
                subset=["ID_Email"],
                keep="last",
            )
        )

        df_sem_id = df_historico[
            ~ids_validos
        ]

        df_historico = pd.concat(
            [
                df_com_id,
                df_sem_id,
            ],
            ignore_index=True,
        )

    # ------------------------------------------------------
    # GERA E SALVA AS VISÕES
    # ------------------------------------------------------

    visoes = gerar_visoes_historico(
        df_historico
    )

    try:
        salvar_visoes_log(
            arquivo_log=arquivo_log,
            visoes=visoes,
        )

    except PermissionError:
        raise

    except Exception as erro:
        print(
            "\nOcorreu um erro ao salvar "
            "o arquivo de log."
        )

        print(
            f"Arquivo: {arquivo_log}"
        )

        print(
            f"Erro: {erro}"
        )

        traceback.print_exc()

        return False

    # ------------------------------------------------------
    # RESUMO DA EXECUÇÃO ATUAL
    # ------------------------------------------------------

    print("\n" + "=" * 80)
    print("RESUMO DA EXECUÇÃO ATUAL")
    print("=" * 80)

    print(
        f"Controle atualizado em: {arquivo_log}"
    )

    print(
        "Fornecedores processados nesta execução: "
        f"{qtd_processados_execucao}"
    )

    print(
        "E-mails enviados nesta execução: "
        f"{qtd_enviados_execucao}"
    )

    print(
        "E-mails não enviados nesta execução: "
        f"{qtd_nao_enviados_execucao}"
    )

    print(
        "E-mails abertos para revisão nesta execução: "
        f"{qtd_revisao_execucao}"
    )

    print(
        "Aguardando retorno nesta execução: "
        f"{qtd_aguardando_execucao}"
    )

    print("=" * 80)

    return True


# ==========================================================
# FECHAMENTO DO STATUS DE RETORNO
# ==========================================================

def fechar_status_retorno(
    arquivo_log,
    id_email,
    status_retorno,
    data_resposta=None,
    email_resposta="",
    assunto_resposta="",
    arquivo_resposta="",
    motivo_resposta="",
):
    """
    Fecha ou atualiza o status de retorno de um e-mail.

    A alteração é realizada na aba Historico, que é a
    fonte oficial do acompanhamento.

    Após a alteração, todas as abas derivadas são
    recriadas automaticamente.

    Retorna:
        True quando o ID for atualizado.
        False quando o ID não for localizado.
    """

    arquivo_log = Path(
        arquivo_log
    )

    if not arquivo_log.exists():
        raise FileNotFoundError(
            "Arquivo de acompanhamento não encontrado: "
            f"{arquivo_log}"
        )

    id_email = str(
        id_email
    ).strip().upper()

    status_retorno = str(
        status_retorno
    ).strip().upper()

    if not id_email:
        raise ValueError(
            "O ID do e-mail precisa ser informado."
        )

    if not status_retorno:
        raise ValueError(
            "O status de retorno precisa ser informado."
        )

    # ------------------------------------------------------
    # LEITURA DO HISTÓRICO
    # ------------------------------------------------------

    try:
        df_historico = pd.read_excel(
            arquivo_log,
            sheet_name="Historico",
            engine="openpyxl",
        )

    except PermissionError as erro:
        raise PermissionError(
            "Não foi possível abrir o arquivo de "
            "acompanhamento. Feche o arquivo no Excel "
            "e tente novamente."
        ) from erro

    except ValueError as erro:
        raise ValueError(
            "A aba 'Historico' não foi encontrada no "
            "arquivo de acompanhamento."
        ) from erro

    except Exception as erro:
        raise RuntimeError(
            "Não foi possível ler o histórico: "
            f"{erro}"
        ) from erro

    # ------------------------------------------------------
    # GARANTE AS COLUNAS
    # ------------------------------------------------------

    df_historico = garantir_colunas_log(
        df_historico
    )

    colunas_texto_retorno = [
        "StatusRetorno",
        "EmailResposta",
        "AssuntoResposta",
        "ArquivoResposta",
        "MotivoResposta",
    ]

    for coluna in colunas_texto_retorno:
        df_historico[coluna] = (
            df_historico[coluna]
            .fillna("")
            .astype("object")
        )

    # ------------------------------------------------------
    # CORRIGE O TIPO DA DATARESPOSTA
    # ------------------------------------------------------

    # Quando uma coluna está totalmente vazia no Excel,
    # o pandas pode carregá-la como float64.
    #
    # A conversão explícita abaixo transforma a coluna em
    # datetime64[ns], permitindo receber uma data.
    serie_data_resposta = pd.to_datetime(
        df_historico["DataResposta"],
        errors="coerce",
        utc=True,
    )

    df_historico["DataResposta"] = (
        serie_data_resposta.dt.tz_localize(None)
        .astype("datetime64[ns]")
    )

    # ------------------------------------------------------
    # LOCALIZA O ID
    # ------------------------------------------------------

    ids_historico = (
        df_historico["ID_Email"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )

    mascara_id = ids_historico.eq(
        id_email
    )

    if not mascara_id.any():
        print(
            f"ID não encontrado no histórico: {id_email}"
        )

        return False

    # ------------------------------------------------------
    # CONVERTE A DATA RECEBIDA DO OUTLOOK
    # ------------------------------------------------------

    data_resposta_convertida = None

    if data_resposta is not None:
        try:
            # O Outlook pode retornar um datetime com um objeto
            # de timezone incompatível com o pandas.
            #
            # Por isso, criamos um novo datetime sem timezone,
            # preservando exatamente os componentes da data.
            data_resposta_convertida = datetime(
                year=data_resposta.year,
                month=data_resposta.month,
                day=data_resposta.day,
                hour=data_resposta.hour,
                minute=data_resposta.minute,
                second=data_resposta.second,
                microsecond=data_resposta.microsecond,
            )

            data_resposta_convertida = (
                pd.Timestamp(data_resposta_convertida)
                .floor("s")
                .as_unit("ns")
            )

        except Exception as erro:
            raise ValueError(
                "Não foi possível converter a data "
                f"da resposta: {data_resposta}"
            ) from erro
    # ------------------------------------------------------
    # ATUALIZA O REGISTRO
    # ------------------------------------------------------

    df_historico.loc[
        mascara_id,
        "StatusRetorno",
    ] = status_retorno

    if data_resposta_convertida is not None:
        df_historico.loc[
            mascara_id,
            "DataResposta",
        ] = data_resposta_convertida

    if email_resposta:
        df_historico.loc[
            mascara_id,
            "EmailResposta",
        ] = str(
            email_resposta
        ).strip()

    if assunto_resposta:
        df_historico.loc[
            mascara_id,
            "AssuntoResposta",
        ] = str(
            assunto_resposta
        ).strip()

    if arquivo_resposta:
        df_historico.loc[
            mascara_id,
            "ArquivoResposta",
        ] = str(
            arquivo_resposta
        ).strip()

    if motivo_resposta:
        df_historico.loc[
            mascara_id,
            "MotivoResposta",
        ] = str(
            motivo_resposta
        ).strip()

    # ------------------------------------------------------
    # GARANTE DATAS SEM TIMEZONE
    # ------------------------------------------------------

    df_historico = converter_colunas_data(
        df_historico
    )

    # ------------------------------------------------------
    # RECRIA AS ABAS DERIVADAS
    # ------------------------------------------------------

    visoes = gerar_visoes_historico(
        df_historico
    )

    # ------------------------------------------------------
    # SALVA O ARQUIVO
    # ------------------------------------------------------

    try:
        salvar_visoes_log(
            arquivo_log=arquivo_log,
            visoes=visoes,
        )

    except PermissionError:
        raise

    except Exception as erro:
        raise RuntimeError(
            "Não foi possível salvar o fechamento "
            f"do retorno do ID {id_email}: {erro}"
        ) from erro

    print(
        f"Status do ID {id_email} atualizado para "
        f"'{status_retorno}'."
    )

    return True