from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def destacar_coluna_excel(
    caminho_arquivo,
    nome_coluna,
):
    """
    Destaca uma coluna específica do Excel.

    Formatação:
    - cabeçalho vermelho com fonte branca;
    - células da coluna com fundo vermelho claro;
    - fonte preta;
    - alinhamento central;
    - quebra automática de texto.
    """

    caminho_arquivo = Path(caminho_arquivo)

    if not caminho_arquivo.exists():
        raise FileNotFoundError(
            f"Arquivo não encontrado: {caminho_arquivo}"
        )

    if not nome_coluna or not str(nome_coluna).strip():
        raise ValueError(
            "O nome da coluna precisa ser informado."
        )

    workbook = None

    try:
        workbook = load_workbook(caminho_arquivo)

        aba = workbook.active

        nome_procurado = str(nome_coluna).strip().casefold()

        coluna_encontrada = None
        celula_cabecalho = None

        # Procura somente na primeira linha.
        for celula in aba[1]:
            
            if celula.value is None:
                continue

            nome_cabecalho = str(celula.value).strip().casefold()

            if nome_cabecalho == nome_procurado:
                coluna_encontrada = celula.column
                celula_cabecalho = celula
                break

        if coluna_encontrada is None:
            cabecalhos = [
                str(celula.value).strip()
                for celula in aba[1]
                if celula.value is not None
            ]

            raise KeyError(
                f"Coluna '{nome_coluna}' não encontrada. "
                f"Colunas disponíveis: {cabecalhos}"
            )

        # Cores com 8 caracteres no padrão ARGB.
        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            start_color="FFED1C24",
            end_color="FFED1C24",
        )

        preenchimento_dados = PatternFill(
            fill_type="solid",
            start_color="FFFFC7CE",
            end_color="FFFFC7CE",
        )

        fonte_cabecalho = Font(
            color="FFFFFFFF",
            bold=True,
        )

        fonte_dados = Font(
            color="FF000000",
            bold=False,
        )

        alinhamento = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        # Formata o cabeçalho.
        celula_cabecalho.fill = preenchimento_cabecalho
        celula_cabecalho.font = fonte_cabecalho
        celula_cabecalho.alignment = alinhamento

        # Formata todas as linhas da coluna.
        for numero_linha in range(2, aba.max_row + 1):
            celula = aba.cell(
                row=numero_linha,
                column=coluna_encontrada,
            )

            celula.fill = preenchimento_dados
            celula.font = fonte_dados
            celula.alignment = alinhamento

        # Ajusta a largura da coluna.
        letra_coluna = celula_cabecalho.column_letter
        aba.column_dimensions[letra_coluna].width = 35

        workbook.save(caminho_arquivo)

        print(
            f"Coluna destacada com sucesso: "
            f"{nome_coluna} | "
            f"coluna Excel: {letra_coluna} | "
            f"posição: {coluna_encontrada}"
        )

        return caminho_arquivo

    except PermissionError as erro:
        raise PermissionError(
            f"Não foi possível salvar o arquivo "
            f"'{caminho_arquivo.name}'. "
            "Verifique se o arquivo está aberto no Excel."
        ) from erro

    finally:
        if workbook is not None:
            workbook.close()